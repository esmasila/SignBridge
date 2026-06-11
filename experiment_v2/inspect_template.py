"""Sablonu detayli incele - hangi shape nerede, ne icerik var."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Emu

PATH = r"C:\Users\leven\Downloads\template.pptx"
prs = Presentation(PATH)

def emu_to_cm(emu):
    return round(emu / 360000, 2)

print(f"Slide boyutu: {emu_to_cm(prs.slide_width)} x {emu_to_cm(prs.slide_height)} cm")
print(f"Slide sayisi: {len(prs.slides)}")

for s_idx, slide in enumerate(prs.slides):
    print(f"\n{'='*70}\nSLIDE {s_idx+1}\n{'='*70}")
    for i, sh in enumerate(slide.shapes):
        x, y = emu_to_cm(sh.left), emu_to_cm(sh.top)
        w, h = emu_to_cm(sh.width), emu_to_cm(sh.height)
        kind = sh.shape_type
        name = sh.name
        text = ""
        if sh.has_text_frame:
            text = sh.text_frame.text.strip()[:80].replace("\n", " | ")
        is_pic = sh.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE
        print(f"[{i:2d}] {name:25s}  ({x:.1f},{y:.1f}) {w:.1f}x{h:.1f}  type={kind}  pic={is_pic}")
        if text:
            print(f"      text: \"{text}\"")
