# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(12)

def h(t, l=1):
    p = doc.add_heading(t, level=l)
    p.style.font.color.rgb = RGBColor(0x5B, 0x21, 0xB6)
    return p

def para(t, italic=False):
    pr = doc.add_paragraph(t)
    if italic:
        pr.runs[0].italic = True
    return pr

title = doc.add_heading('SignBridge — Tez Sunum Konuşması', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
s = doc.add_paragraph('Teknik terime girmeden, anlaşılır anlatım')
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
s.runs[0].font.color.rgb = RGBColor(0x64,0x74,0x8B)
s.runs[0].italic = True
doc.add_paragraph()

# === DIZI ACIKLAMASI ===
h('"Dizi" Nedir? (Kısa Açıklama)', 1)
para('İşaret dili hareket içerir — tek bir fotoğraf bir kelimeyi anlatmaya yetmez. '
     'Bu yüzden bir kelime işaret edildiğinde kamera yaklaşık 1 saniye boyunca 30 ardışık '
     'kare çeker. Bu 30 kare birlikte bir "dizi" oluşturur. Yani bir dizi = bir kelimenin '
     'bir kez gösterimi = hareketin film şeridi gibi kaydı.')

# === ACILIS ===
h('1. Açılış', 1)
para('"Merhaba sayın hocalarım. Bugün sizlere bitirme projem SignBridge\'i tanıtacağım.')
para('')
para('Türkiye\'de yaklaşık 700 bin işitme engelli vatandaşımız var. Onların günlük hayatta '
     'iletişim kurması çoğu zaman zor — çünkü etrafındaki insanların büyük kısmı işaret dili '
     'bilmiyor. İşaret dili öğrenmek zaman alıyor, tercüman bulmak ise pahalı ve her an mümkün değil.')
para('')
para('Ben de düşündüm: Acaba teknoloji bu iletişim engelini ortadan kaldırabilir mi? '
     'İşte SignBridge bu sorudan doğdu."')

# === PROJE NEDIR ===
h('2. Proje Nedir?', 1)
para('"SignBridge, Türk İşaret Dili ile Türkçe arasında çeviri yapan bir uygulamadır. '
     'En önemli özelliği çift yönlü çalışmasıdır:')
para('')
para('Birinci yön: İşitme engelli bir kişi kameraya işaret yapar, uygulama bunu anında '
     'Türkçe yazıya ve sese çevirir. Böylece karşısındaki kişi ne dediğini anlar.')
para('')
para('İkinci yön: Konuşan bir kişi Türkçe yazar ya da mikrofona konuşur, ekrandaki '
     '3 boyutlu avatar bunu işaret diliyle gösterir. Böylece işitme engelli kişi de anlar.')
para('')
para('Yani SignBridge iki tarafın da birbirini anlamasını sağlayan bir köprü — ismi de '
     'buradan geliyor."')

# === NASIL CALISIYOR ===
h('3. Nasıl Çalışıyor? (Basit Anlatım)', 1)
para('"Sistemi çok teknik anlatmadan özetleyeyim:')
para('')
para('Kullanıcı kameraya işaret yaptığında, uygulama elin ve vücudun hareketini takip ediyor. '
     'Daha önceden binlerce örnekle eğittiğim yapay zeka modeli, bu hareketi tanıyıp hangi '
     'kelime olduğunu buluyor. Sonra bu kelimeleri anlamlı bir Türkçe cümleye dönüştürüyor '
     've sesli olarak okuyor.')
para('')
para('Avatar tarafında ise tam tersi oluyor: Yazılan cümle önce kelimelere ayrılıyor, sonra '
     'her kelime için önceden hazırladığım hareketi 3 boyutlu karakter sırayla yapıyor."')

# === VERI SETI ===
h('4. Veri Seti', 1)
para('"Yapay zekanın bir şeyi tanıması için ona bolca örnek göstermek gerekiyor.')
para('')
para('Ben 122 kelimelik bir sözlük oluşturdum — günlük hayatta en çok kullanılan kelimeler: '
     'aile bireyleri, sayılar, günler, selamlaşma sözleri, sık kullanılan fiiller.')
para('')
para('Her kelimeyi kameranın önünde defalarca işaret ederek kaydettim. Toplamda yaklaşık '
     '15 bin örnek topladım. Her örnek, bir kelimenin yaklaşık 1 saniyelik kaydı.')
para('')
para('Sonra bu veriyi yapay olarak çoğalttım — her kaydı hafifçe değiştirip yeni varyantlar '
     'ürettim. Böylece veri setim yaklaşık 105 bin örneğe ulaştı. Bu, modelin farklı '
     'koşullarda da iyi çalışmasını sağladı.')
para('')
para('Önemli bir nokta: Veri setimde fotoğraf saklamıyorum. Sadece elin ve vücudun '
     'koordinat noktalarını tutuyorum. Bu hem dosyaları küçük tutuyor hem de kişinin '
     'görüntüsü kaydedilmediği için gizliliği koruyor."')

# === SONUC ===
h('5. Sonuç ve Hedefler', 1)
para('"SignBridge şu an 122 kelimeyi tanıyabiliyor, hem bilgisayarda hem de telefonda '
     'çalışıyor. İnternet olmadan, tamamen cihaz üzerinde de çalışabiliyor.')
para('')
para('Gelecekte sözlüğü genişletmek, cümle yapısını daha akıcı hale getirmek ve gerçek '
     'işitme engelli bireylerle test etmek istiyorum.')
para('')
para('SignBridge\'in amacı sadece bir bitirme projesi olmak değil — gerçekten işe yarayan, '
     'insanların hayatını kolaylaştıran bir araç olmasını hedefliyorum.')
para('')
para('Beni dinlediğiniz için teşekkür ederim. Sorularınızı almaktan memnuniyet duyarım."')

# === DEMO ===
h('6. Canlı Demo Anı (Söyleyecekler)', 1)
para('"Şimdi izninizle sistemi canlı göstereyim.')
para('')
para('[Kamera sekmesinde işaret yap] Şu an kameraya bir işaret yapıyorum... gördüğünüz '
     'gibi uygulama kelimeyi tanıdı ve ekrana yazdı, sesli de okudu.')
para('')
para('[Avatar sekmesine geç] Şimdi de tersini göstereyim. Buraya bir cümle yazıyorum... '
     've avatar bunu işaret diliyle gösteriyor.')
para('')
para('[Hızlı cümleler] Hazır cümlelerden de seçebiliyorum, böylece günlük bir diyalog '
     'akışını canlı izleyebilirsiniz."')

doc.add_paragraph()
e = doc.add_paragraph('İpucu: Konuşurken acele etme, demo sırasında "konuş ve göster" — '
                      'ekran sessiz kalmasın. Toplam süre 7-9 dakika ideal.')
e.runs[0].italic = True
e.runs[0].font.color.rgb = RGBColor(0x64,0x74,0x8B)

out = r'C:\Users\leven\Downloads\SignBridge-Sunum-Konusmasi.docx'
doc.save(out)
print('Olusturuldu:', out)
